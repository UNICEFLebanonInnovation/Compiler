"""Regression coverage for registration photographs embedded in MSCC Excel exports."""

import csv
import io
import json
import posixpath
import tempfile
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from urllib.parse import urlparse
from xml.etree import ElementTree

from django.contrib.auth import get_user_model
from django.contrib.auth.models import AnonymousUser, Group
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.storage import FileSystemStorage
from django.http import Http404, HttpResponse
from django.test import RequestFactory, SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from student_registration.backends.models import ExportHistory
from student_registration.child.models import Child
from student_registration.locations.models import Center
from student_registration.mscc.document_exports import (
    PHOTO_INCLUDED, PHOTO_MISSING, PHOTO_UNAVAILABLE, build_children_workbook,
    export_columns, load_registration_photo,
)
from student_registration.mscc.models import FollowUpService, Registration, Round
from student_registration.schools.models import PartnerOrganization
from student_registration.students.models import Nationality


def picture_bytes(colour, size=(18, 24)):
    output = io.BytesIO()
    Image.new('RGB', size, colour).save(output, format='PNG')
    return output.getvalue()


def workbook_sheet(data):
    return load_workbook(io.BytesIO(data)).active


def workbook_photos(data):
    """Read images and their row anchors directly from the saved OOXML.

    Compiler pins openpyxl 2.4, whose reader does not load worksheet images.
    Inspect the exported document itself rather than depending on its reader.
    """
    drawing_ns = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
    picture_ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    relationship_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    photos = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        drawing_files = sorted(name for name in archive.namelist() if
                               name.startswith('xl/drawings/drawing') and name.endswith('.xml'))
        for filename in drawing_files:
            directory, name = posixpath.split(filename)
            relationships = ElementTree.fromstring(archive.read(
                '{}/_rels/{}.rels'.format(directory, name),
            ))
            targets = {relationship.attrib['Id']: posixpath.normpath(posixpath.join(
                directory, relationship.attrib['Target'],
            )) for relationship in relationships}
            for anchor in ElementTree.fromstring(archive.read(filename)):
                blip = anchor.find('.//{{{}}}blip'.format(picture_ns))
                if blip is None:
                    continue
                start = anchor.find('{{{}}}from'.format(drawing_ns))
                photos.append(SimpleNamespace(
                    row=int(start.find('{{{}}}row'.format(drawing_ns)).text),
                    column=int(start.find('{{{}}}col'.format(drawing_ns)).text),
                    data=archive.read(targets[blip.attrib['{{{}}}embed'.format(relationship_ns)]]),
                ))
    return photos


def image_colour(image):
    with Image.open(io.BytesIO(image.data)) as photograph:
        return photograph.convert('RGB').getpixel((0, 0))


class LocalPhotoStorageMixin:
    def setUp(self):
        super().setUp()
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = directory.name
        settings = override_settings(
            DEFAULT_FILE_STORAGE='django.core.files.storage.FileSystemStorage',
            MEDIA_ROOT=self.directory,
        )
        settings.enable()
        self.addCleanup(settings.disable)
        self.storage = FileSystemStorage(location=self.directory)

    def save_photo(self, registration, colour, name='photo.png', save=False):
        registration.profile_picture.save(name, ContentFile(picture_bytes(colour)), save=save)


class PhotoWorkbookTests(LocalPhotoStorageMixin, SimpleTestCase):
    def registration(self, identifier=101):
        child = Child(
            pk=22, first_name='آدم', father_name='بسام', last_name='العلي',
            mother_fullname='أمل', birthday_day='5', birthday_month='6',
            birthday_year='2016',
        )
        return Registration(pk=identifier, child=child)

    def test_two_registrations_for_same_child_keep_their_own_photos_and_row_anchors(self):
        first = self.registration(101)
        second = self.registration(102)
        self.save_photo(first, 'red', 'first.png')
        self.save_photo(second, 'blue', 'second.png')

        data = build_children_workbook([first, second], fields=['first_name'])
        sheet = workbook_sheet(data)
        self.assertEqual([tuple(value if value != '' else None for value in row)
                          for row in sheet.values], [
            ('Registration ID', 'Child ID', 'Photograph', 'Photograph status', 'First name'),
            ('101', '22', None, PHOTO_INCLUDED, 'آدم'),
            ('102', '22', None, PHOTO_INCLUDED, 'آدم'),
        ])
        photos = workbook_photos(data)
        self.assertEqual([(image.row, image.column) for image in photos], [(1, 2), (2, 2)])
        self.assertEqual([image_colour(image) for image in photos],
                         [(255, 0, 0), (0, 0, 255)])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            drawing = ElementTree.fromstring(archive.read('xl/drawings/drawing1.xml'))
        # The two-cell anchor moves and sizes with its row when Excel filters it.
        self.assertTrue(all(anchor.tag.endswith('twoCellAnchor') for anchor in drawing))
        self.assertTrue(all(anchor.attrib.get('editAs', 'twoCell') == 'twoCell'
                            for anchor in drawing))

    def test_replaced_photo_is_read_from_current_registration_storage(self):
        registration = self.registration()
        self.save_photo(registration, 'red', 'old.png')
        self.assertEqual(image_colour(workbook_photos(
            build_children_workbook([registration]))[0]), (255, 0, 0))
        self.save_photo(registration, 'green', 'replacement.png')
        self.assertEqual(image_colour(workbook_photos(
            build_children_workbook([registration]))[0]), (0, 128, 0))

    def test_missing_corrupt_and_unreadable_photos_keep_registration_rows(self):
        missing = self.registration(101)
        corrupt = self.registration(102)
        corrupt.profile_picture.save('corrupt.png', ContentFile(b'not an image'), save=False)
        unreadable = self.registration(103)
        unreadable.profile_picture = 'mscc/profile_pictures/missing.png'
        data = build_children_workbook([missing, corrupt, unreadable], fields=['first_name'])
        sheet = workbook_sheet(data)
        self.assertEqual(sheet.max_row, 4)
        self.assertFalse(workbook_photos(data))
        self.assertEqual([sheet.cell(row=row, column=4).value for row in range(2, 5)],
                         [PHOTO_MISSING, PHOTO_UNAVAILABLE, PHOTO_UNAVAILABLE])

        missing.child.photo = Mock(name='unrelated_child_photo')
        self.assertEqual(load_registration_photo(missing), (None, PHOTO_MISSING))
        missing.child.photo.open.assert_not_called()

    def test_remote_storage_error_is_reported_without_storage_path_access(self):
        registration = self.registration()
        registration.profile_picture = 'remote/photo.jpg'
        remote = Mock()
        remote.open.side_effect = OSError('storage unavailable')
        registration.profile_picture.storage = remote
        self.assertEqual(load_registration_photo(registration), (None, PHOTO_UNAVAILABLE))
        remote.open.assert_called_once_with('remote/photo.jpg', 'rb')
        remote.path.assert_not_called()

    def test_exif_orientation_is_applied_and_metadata_is_removed(self):
        registration = self.registration()
        output = io.BytesIO()
        image = Image.new('RGB', (12, 8), 'red')
        exif = image.getexif()
        exif[274] = 6
        exif[270] = 'private camera metadata'
        image.save(output, format='JPEG', exif=exif)
        registration.profile_picture.save('oriented.jpg', ContentFile(output.getvalue()), save=False)

        data, status = load_registration_photo(registration)
        self.assertEqual(status, PHOTO_INCLUDED)
        with Image.open(io.BytesIO(data)) as photograph:
            self.assertEqual(photograph.format, 'PNG')
            self.assertEqual(photograph.size, (8, 12))
            self.assertFalse(photograph.getexif())
            self.assertNotIn('private camera metadata', str(photograph.info))

    def test_selected_columns_keep_identity_and_names_are_literal_xml_safe_text(self):
        registration = self.registration()
        registration.child.first_name = '=HYPERLINK("https://example.invalid", "name")\x00\ud800'
        registration.child.last_name = 'https://example.invalid'
        sheet = workbook_sheet(build_children_workbook(
            [registration], fields=['first_name', 'last_name'],
        ))
        self.assertEqual(sheet.max_column, 6)
        self.assertEqual(sheet.cell(row=2, column=1).value, '101')
        self.assertEqual(sheet.cell(row=2, column=2).value, '22')
        self.assertEqual(sheet.cell(row=2, column=5).value,
                         '=HYPERLINK("https://example.invalid", "name")')
        self.assertEqual(sheet.cell(row=2, column=5).data_type, 's')
        self.assertIsNone(sheet.cell(row=2, column=6).hyperlink)
        self.assertEqual([column.key for column in export_columns(['unknown'])],
                         ['registration_id', 'child_id', 'photograph', 'photograph_status'])

    def test_https_middleware_preserves_private_photo_export_cache_policy(self):
        from student_registration.cache_control_middleware import CacheControlMiddleware

        request = RequestFactory().get('/mscc/export-download/photo.zip/', secure=True)
        response = HttpResponse(b'private photo archive')
        response['Cache-Control'] = 'private, no-store'
        middleware = CacheControlMiddleware(lambda incoming: response)
        result = middleware(request)
        self.assertEqual(result['Cache-Control'], 'private, no-store')


class PhotoExcelExportTests(LocalPhotoStorageMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.factory = RequestFactory()
        self.partner = PartnerOrganization.objects.create(name='Photo export partner')
        self.outside_partner = PartnerOrganization.objects.create(name='Other photo export partner')
        self.center = Center.objects.create(name='Photo export center', partner=self.partner)
        self.sibling_center = Center.objects.create(name='Sibling center', partner=self.partner)
        self.outside_center = Center.objects.create(name='Outside center', partner=self.outside_partner)
        self.round = Round.objects.create(name='Photo export round')
        self.nationality = Nationality.objects.create(name='Photo export nationality')
        self.child = Child.objects.create(
            first_name='آدم', last_name='العلي', father_name='بسام', mother_fullname='أمل',
            birthday_day='5', birthday_month='6', birthday_year='2016', gender='Male',
            nationality=self.nationality,
        )
        self.user = self.make_user('photo-center', 'MSCC_CENTER', self.partner, self.center)
        self.first = self.make_registration(self.center, self.partner, self.round)
        self.second = self.make_registration(self.center, self.partner, None)
        self.sibling = self.make_registration(self.sibling_center, self.partner, self.round)
        self.outside = self.make_registration(self.outside_center, self.outside_partner, self.round)
        self.make_registration(self.center, self.partner, self.round, deleted=True)
        self.make_registration(self.center, self.partner, self.round, type='TLS')
        Registration.objects.create(child=None, center=self.center, partner=self.partner)
        self.save_photo(self.first, 'red', 'registered-red.png', save=True)
        self.save_photo(self.second, 'blue', 'registered-blue.png', save=True)

    def make_user(self, username, group=None, partner=None, center=None, **kwargs):
        user = get_user_model().objects.create_user(
            username=username, partner=partner, center=center, **kwargs,
        )
        if group:
            user.groups.add(Group.objects.get_or_create(name=group)[0])
        return user

    def make_registration(self, center, partner, cycle, **kwargs):
        return Registration.objects.create(
            child=self.child, center=center, partner=partner, round=cycle, **kwargs,
        )

    def make_export(self, user=None, **kwargs):
        return ExportHistory.objects.create(
            created_by=user or self.user, export_type='Makani List', file_format='xlsx', **kwargs,
        )

    def request(self, method, url, data=None, user=None):
        if method == 'post':
            request = self.factory.post(url, json.dumps(data or {}), content_type='application/json')
        else:
            request = self.factory.get(url, data or {})
        request.user = user or self.user
        return request

    def test_export_roles_scope_to_accessible_registered_children(self):
        from student_registration.mscc.export_access import excel_export_registrations, user_can_export

        own = {self.first.pk, self.second.pk}
        partner = own | {self.sibling.pk}
        everyone = partner | {self.outside.pk}
        global_user = self.make_user('photo-unicef', 'MSCC_UNICEF')
        partner_user = self.make_user('photo-partner', 'MSCC_PARTNER', self.partner)
        wl_partner = PartnerOrganization.objects.create(name='World Learning', is_world_learning=True)
        world_learning = self.make_user('photo-world-learning', partner=wl_partner)
        staff = self.make_user('photo-staff', is_staff=True)
        admin = self.make_user('photo-admin', is_superuser=True)
        inactive = self.make_user('photo-inactive', 'MSCC_UNICEF', is_active=False)
        unassigned = self.make_user('photo-unassigned', 'MSCC_CENTER')
        unrelated = self.make_user('photo-unrelated')
        for user, expected in [
            (self.user, own), (partner_user, partner), (global_user, everyone),
            (world_learning, everyone), (staff, everyone), (admin, everyone),
            (inactive, set()), (unassigned, set()), (unrelated, set()),
            (AnonymousUser(), set()),
        ]:
            with self.subTest(user=str(user)):
                self.assertEqual(user_can_export(user), bool(expected))
                self.assertEqual(set(excel_export_registrations(user).values_list('pk', flat=True)), expected)

    def test_round_and_identity_filters_are_applied_inside_the_role_scope(self):
        from student_registration.mscc.export_access import excel_export_registrations

        for filters, expected in [
            ({}, {self.first.pk, self.second.pk}),
            ({'round': str(self.round.pk)}, {self.first.pk}),
            ({'round': 'no_round'}, {self.second.pk}),
            ({'nationality': str(self.nationality.pk), 'first_name': 'آد',
              'last_name': 'علي', 'father_name': 'بسا', 'mother_fullname': 'أمل'},
             {self.first.pk, self.second.pk}),
            ({'first_name': 'someone else'}, set()),
        ]:
            with self.subTest(filters=filters):
                self.assertEqual(set(excel_export_registrations(
                    self.user, **filters).values_list('pk', flat=True)), expected)
        for filters in [{'round': 'invalid'}, {'round': '-1'}, {'nationality': 'invalid'}]:
            with self.subTest(filters=filters), self.assertRaises(ValidationError):
                excel_export_registrations(self.user, **filters)

    def test_worker_generates_real_photo_zip_and_notification_failure_preserves_download(self):
        from student_registration.mscc.tasks import _generate_excel_mscc_export

        export = self.make_export()
        with patch('student_registration.mscc.tasks.ExportStorage', return_value=self.storage), patch(
                'student_registration.mscc.tasks.send_push_to_web', side_effect=RuntimeError('push unavailable')):
            _generate_excel_mscc_export(export.pk, fields=['first_name'])
        export.refresh_from_db()
        self.assertEqual(export.status, 'done')
        filename = Path(urlparse(export.file_url).path.rstrip('/')).name
        self.assertTrue(filename.startswith('mscc_photo_export_'))
        with self.storage.open(filename, 'rb') as source, zipfile.ZipFile(source) as archive:
            self.assertEqual(archive.namelist(), ['mscc_data.xlsx'])
            data = archive.read('mscc_data.xlsx')
            sheet = workbook_sheet(data)
        self.assertEqual({sheet.cell(row=row, column=1).value for row in (2, 3)},
                         {str(self.first.pk), str(self.second.pk)})
        self.assertEqual([image_colour(image) for image in workbook_photos(data)],
                         [(255, 0, 0), (0, 0, 255)])

    def test_filtered_excel_contains_scoped_followup_data_and_rechecks_worker_permission(self):
        from student_registration.mscc.tasks import _generate_excel_mscc_export

        FollowUpService.objects.create(registration=self.first, follow_up_type='Phone call')
        FollowUpService.objects.create(registration=self.outside, follow_up_type='Home Visits')
        export = self.make_export()
        with patch('student_registration.mscc.tasks.ExportStorage', return_value=self.storage), patch(
                'student_registration.mscc.tasks.send_push_to_web'):
            _generate_excel_mscc_export(
                export.pk, round=str(self.round.pk), include_followup=True,
            )
        export.refresh_from_db()
        filename = Path(urlparse(export.file_url).path.rstrip('/')).name
        with self.storage.open(filename, 'rb') as source, zipfile.ZipFile(source) as archive:
            self.assertEqual(set(archive.namelist()), {'mscc_data.xlsx', 'followup_data.csv'})
            sheet = workbook_sheet(archive.read('mscc_data.xlsx'))
            followups = list(csv.DictReader(io.StringIO(
                archive.read('followup_data.csv').decode('utf-8-sig'),
            )))
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(row=2, column=1).value, str(self.first.pk))
        self.assertEqual([row['registration_id'] for row in followups], [str(self.first.pk)])

        denied = self.make_export()
        self.user.groups.clear()
        with patch('student_registration.mscc.tasks.ExportStorage') as storage, patch(
                'student_registration.mscc.tasks.send_push_to_web'):
            _generate_excel_mscc_export(denied.pk)
        denied.refresh_from_db()
        self.assertEqual(denied.status, 'failed')
        self.assertFalse(denied.file_url)
        storage.return_value.save.assert_not_called()

    def test_both_export_endpoints_queue_excel_with_saved_format_and_validate_payloads(self):
        from student_registration.mscc import views

        with patch('student_registration.mscc.views.queue_mscc_export') as queue:
            response = views.export_list_async(self.request('post', '/export/', {
                'format': 'xlsx', 'fields': ['first_name'],
            }))
        self.assertEqual(response.status_code, 200)
        export = ExportHistory.objects.get()
        self.assertEqual(export.file_format, 'xlsx')
        self.assertEqual(export.fields, ['first_name'])
        queue.assert_called_once_with(export.pk, ['first_name'], 'xlsx')

        with patch('student_registration.mscc.views.queue_filtered_mscc_export') as queue:
            response = views.export_list_background(self.request('get', '/filtered/', {
                'format': 'xlsx', 'round': str(self.round.pk), 'first_name': 'آدم',
            }))
        self.assertEqual(response.status_code, 200)
        export = ExportHistory.objects.order_by('-pk').first()
        self.assertEqual(export.file_format, 'xlsx')
        self.assertEqual(queue.call_args.args[2], 'آدم')
        self.assertEqual(queue.call_args.kwargs['file_format'], 'xlsx')

        for payload in [{'format': 'pdf'}, {'format': 'xlsx', 'fields': 'first_name'},
                        {'format': 'xlsx', 'fields': [123]}]:
            with self.subTest(payload=payload), patch(
                    'student_registration.mscc.views.queue_mscc_export') as queue:
                before = ExportHistory.objects.count()
                response = views.export_list_async(self.request('post', '/export/', payload))
                self.assertEqual(response.status_code, 400)
                self.assertEqual(ExportHistory.objects.count(), before)
                queue.assert_not_called()

        for raw_payload in ['["xlsx"]', 'null', '123', '"xlsx"', '{invalid json']:
            with self.subTest(raw_payload=raw_payload), patch(
                    'student_registration.mscc.views.queue_mscc_export') as queue:
                before = ExportHistory.objects.count()
                request = self.factory.post('/export/', raw_payload, content_type='application/json')
                request.user = self.user
                response = views.export_list_async(request)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(ExportHistory.objects.count(), before)
                queue.assert_not_called()

    def test_invalid_filtered_requests_and_unassigned_excel_roles_do_not_queue_jobs(self):
        from student_registration.mscc import views

        for parameters in [
            {'format': 'xlsx'}, {'format': 'xlsx', 'round': 'invalid'},
            {'format': 'xlsx', 'round': 'no_round', 'nationality': '-1'},
        ]:
            with self.subTest(parameters=parameters), patch(
                    'student_registration.mscc.views.queue_filtered_mscc_export') as queue:
                response = views.export_list_background(self.request('get', '/filtered/', parameters))
                self.assertEqual(response.status_code, 400)
                queue.assert_not_called()
        unassigned = self.make_user('photo-unassigned-requester')
        with patch('student_registration.mscc.views.queue_mscc_export') as queue:
            response = views.export_list_async(self.request(
                'post', '/export/', {'format': 'xlsx'}, user=unassigned,
            ))
        self.assertEqual(response.status_code, 403)
        queue.assert_not_called()
        with patch('student_registration.mscc.views.queue_filtered_mscc_export') as queue:
            response = views.export_list_background(self.request(
                'get', '/filtered/', {'format': 'xlsx', 'round': str(self.round.pk)}, user=unassigned,
            ))
        self.assertEqual(response.status_code, 403)
        queue.assert_not_called()
        self.assertFalse(ExportHistory.objects.exists())

    def test_photo_export_download_requires_its_creator_and_completed_history(self):
        from student_registration.mscc import views

        filename = 'mscc_photo_export_test.zip'
        export = self.make_export(status='done', file_url=reverse(
            'mscc:export_download', args=[filename],
        ))
        other_user = self.make_user('photo-other-downloader', 'MSCC_UNICEF')
        with patch('student_registration.mscc.views.download_file', return_value=HttpResponse(b'zip')) as download:
            response = views.get_file(self.request('get', '/download/'), filename)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        download.assert_called_once()

        for user, status in [(other_user, 'done'), (self.user, 'pending'), (self.user, 'failed')]:
            with self.subTest(user=user.username, status=status):
                export.status = status
                export.save(update_fields=['status'])
                with patch('student_registration.mscc.views.download_file') as download, self.assertRaises(Http404):
                    views.get_file(self.request('get', '/download/', user=user), filename)
                download.assert_not_called()

    def test_owner_can_download_real_photo_archive_through_shared_storage_helper(self):
        from student_registration.mscc import views

        filename = 'mscc_photo_export_owner.zip'
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as archive:
            archive.writestr('mscc_data.xlsx', build_children_workbook(
                [self.first], fields=['first_name'],
            ))
        self.storage.save(filename, ContentFile(output.getvalue()))
        self.make_export(status='done', file_url=reverse('mscc:export_download', args=[filename]))

        self.client.force_login(self.user)
        with patch('student_registration.backends.utils.ExportStorage', return_value=self.storage):
            response = self.client.get(reverse('mscc:export_download', args=[filename]))
            # Django's client closes the streamed response while preserving
            # the database transaction managed by TestCase.
            downloaded = b''.join(response.streaming_content)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        with zipfile.ZipFile(io.BytesIO(downloaded)) as archive:
            data = archive.read('mscc_data.xlsx')
            sheet = workbook_sheet(data)
        self.assertEqual(sheet.cell(row=2, column=1).value, str(self.first.pk))
        self.assertEqual(image_colour(workbook_photos(data)[0]), (255, 0, 0))

    def test_shared_storage_and_tls_alias_reject_photo_archive_before_storage_access(self):
        from student_registration.backends.utils import download_file
        from student_registration.tls.child_views import get_file as tls_download

        filename = 'mscc_photo_export_private.zip'
        self.make_export(status='done', file_url=reverse('mscc:export_download', args=[filename]))
        other_user = self.make_user('photo-storage-other-user', 'MSCC_UNICEF')
        with patch('student_registration.backends.utils.ExportStorage') as storage:
            with self.assertRaises(Http404):
                download_file(filename, 'export.zip', request=self.request(
                    'get', '/download/', user=other_user,
                ))
            # TLS passes no requester to the shared utility, so even an owner
            # cannot bypass the canonical MSCC route through that older alias.
            with self.assertRaises(Http404):
                tls_download(self.request('get', '/tls/export-download/'), filename)
            storage.assert_not_called()
